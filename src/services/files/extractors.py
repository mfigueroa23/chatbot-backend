"""Texto de documentos de Office y de archivos de texto: Gemini no lee .docx, .xlsx ni .pptx (plan 002, D2)."""
import io
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

def extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        lines += [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
    return "\n".join(lines)

def extract_xlsx(data: bytes) -> str:
    # Valores calculados, no fórmulas: es lo que el colaborador ve en la planilla.
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"Hoja: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value) for value in row]
            if any(cells):
                lines.append(" | ".join(cells).rstrip(" |"))
    workbook.close()
    return "\n".join(lines)

def extract_pptx(data: bytes) -> str:
    presentation = Presentation(io.BytesIO(data))
    lines: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines.append(f"Diapositiva {number}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():  # type: ignore[attr-defined]
                lines.append(shape.text_frame.text)  # type: ignore[attr-defined]
    return "\n".join(lines)

def decode_text(data: bytes) -> str:
    # utf-8-sig quita el BOM que dejan algunos editores; Latin-1 cubre los CSV exportados desde Excel en Windows.
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")
