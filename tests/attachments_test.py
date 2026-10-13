import io
import logging
import pytest
from docx import Document
from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
from langchain_core.messages import AIMessage, BaseMessage
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from src.services import attachments
from src.services.attachments import DOCX, Attachment, read_attachments, resolve_mime
from src.utils.exceptions.attachments import AttachmentTooLargeError


class FakeLLM(GenericFakeChatModel):
    """Modelo falso que guarda los mensajes recibidos: los tests no llaman a Gemini."""
    seen: list[BaseMessage] = []

    def _generate(self, messages, stop=None, run_manager=None, **kwargs):
        self.seen.extend(messages)
        return super()._generate(messages, stop, run_manager, **kwargs)


def fake_llm(*replies: str) -> FakeLLM:
    return FakeLLM(messages=iter([AIMessage(reply) for reply in replies]))


def fake_download(files: dict[str, bytes | Exception]):
    async def download(resource_name: str) -> bytes:
        item = files[resource_name]
        if isinstance(item, Exception):
            raise item
        return item
    return download


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Contrato de crédito automotriz")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Cuota", "250.000"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Pagos"
    sheet.append(["Mes", "Monto"])
    sheet.append(["Octubre", 250000])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Roadmap del portal"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


async def read(items: list[Attachment], files: dict[str, bytes | Exception], llm: FakeLLM | None = None) -> str:
    return await read_attachments(items, fake_download(files), llm or fake_llm())


@pytest.mark.anyio
async def test_lee_word_excel_powerpoint_y_texto_en_memoria():
    files: dict[str, bytes | Exception] = {
        "r/doc": docx_bytes(), "r/xls": xlsx_bytes(), "r/ppt": pptx_bytes(),
        "r/csv": "año;cuota".encode("latin-1"), "r/txt": "﻿cañón".encode("utf-8")}
    items = [Attachment("contrato.docx", DOCX, "r/doc"), Attachment("pagos.xlsx", "application/octet-stream", "r/xls"),
             Attachment("roadmap.pptx", "application/octet-stream", "r/ppt"), Attachment("pagos.csv", "text/csv", "r/csv"),
             Attachment("notas.txt", "text/plain", "r/txt")]

    blocks = (await read(items, files)).split("\n\n")

    assert blocks[0] == "[Archivo adjunto: contrato.docx]\nContrato de crédito automotriz\nCuota | 250.000"
    assert blocks[1] == "[Archivo adjunto: pagos.xlsx]\nHoja: Pagos\nMes | Monto\nOctubre | 250000"
    assert blocks[2] == "[Archivo adjunto: roadmap.pptx]\nDiapositiva 1\nRoadmap del portal"
    assert blocks[3] == "[Archivo adjunto: pagos.csv]\naño;cuota"
    assert blocks[4] == "[Archivo adjunto: notas.txt]\ncañón"


@pytest.mark.parametrize("name, content_type, expected", [
    ("foto.JPG", "image/jpeg", "image/jpeg"),
    ("datos.csv", "application/octet-stream", "text/csv"),
    ("informe.docx", "application/octet-stream", DOCX),
    ("video.mp4", "video/mp4", "video/mp4"),
])
def test_mime_por_tipo_o_por_extension(name, content_type, expected):
    assert resolve_mime(name, content_type) == expected


@pytest.mark.anyio
async def test_imagen_y_pdf_se_leen_con_el_modelo():
    llm = fake_llm("Captura: error 504 en el portal", "Contrato firmado")
    items = [Attachment("foto.png", "image/png", "r/foto"), Attachment("contrato.pdf", "application/pdf", "r/pdf")]

    text = await read(items, {"r/foto": b"png", "r/pdf": b"pdf"}, llm)

    assert "[Archivo adjunto: foto.png]\nCaptura: error 504 en el portal" in text
    assert "[Archivo adjunto: contrato.pdf]\nContrato firmado" in text
    blocks = [message.content[1] for message in llm.seen]
    assert {"type": "image", "base64": "cG5n", "mime_type": "image/png"} in blocks
    assert {"type": "file", "base64": "cGRm", "mime_type": "application/pdf"} in blocks


@pytest.mark.anyio
async def test_un_archivo_con_problemas_no_detiene_a_los_demas(monkeypatch, caplog):
    monkeypatch.setattr(attachments, "MAX_CHARS", 50)
    files: dict[str, bytes | Exception] = {
        "r/grande": AttachmentTooLargeError("r/grande"), "r/roto": ConnectionError("contenido secreto"),
        "r/largo": b"a" * 80, "r/ok": b"hola"}
    items = [Attachment("video.mp4", "video/mp4", "r/video"), Attachment("enlace.pdf", "application/pdf"),
             Attachment("informe.pdf", "application/pdf", "r/grande"), Attachment("roto.txt", "text/plain", "r/roto"),
             Attachment("largo.txt", "text/plain", "r/largo"), Attachment("ok.txt", "text/plain", "r/ok")]

    with caplog.at_level(logging.WARNING):
        blocks = (await read(items, files)).split("\n\n")

    assert blocks[0].startswith("[Archivo adjunto: video.mp4 · formato no admitido: se admiten JPG, PNG")
    assert blocks[1] == "[Archivo adjunto: enlace.pdf · es un enlace de Drive, solo se leen los archivos subidos al chat]"
    assert blocks[2] == "[Archivo adjunto: informe.pdf · supera los 20 MB, no se leyó]"
    assert blocks[3] == "[Archivo adjunto: roto.txt · no se pudo leer]"
    assert blocks[4] == "[Archivo adjunto: largo.txt · leído en parte, es muy largo]\n" + "a" * 50
    assert blocks[5] == "[Archivo adjunto: ok.txt]\nhola"
    assert "roto.txt" in caplog.text and "contenido secreto" not in caplog.text
