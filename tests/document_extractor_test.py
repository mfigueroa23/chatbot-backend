import io
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from src.services.document_extractor import decode_text, extract_docx, extract_pptx, extract_xlsx


def docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Política de viáticos", level=1)
    document.add_paragraph("El tope diario es de 30.000 pesos.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Ciudad", "Tope"
    table.cell(1, 0).text, table.cell(1, 1).text = "Santiago", "30.000"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Cuotas"
    sheet.append(["Mes", "Monto"])
    sheet.append(["Octubre", 150000])
    other = workbook.create_sheet("Notas")
    other.append(["Pagar antes del 5"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Roadmap 2027"  # type: ignore[union-attr]
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "Portal de autoatención"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_extract_docx_incluye_parrafos_y_tablas():
    text = extract_docx(docx_bytes())

    assert "Política de viáticos" in text and "El tope diario es de 30.000 pesos." in text
    assert "Santiago | 30.000" in text


def test_extract_xlsx_incluye_cada_hoja_con_sus_filas():
    text = extract_xlsx(xlsx_bytes())

    assert "Hoja: Cuotas" in text and "Octubre | 150000" in text
    assert "Hoja: Notas" in text and "Pagar antes del 5" in text


def test_extract_pptx_incluye_cada_diapositiva():
    text = extract_pptx(pptx_bytes())

    assert "Diapositiva 1" in text and "Roadmap 2027" in text and "Portal de autoatención" in text


def test_decode_text_utf8_con_bom_y_latin1():
    assert decode_text("﻿cañón,día".encode("utf-8")) == "cañón,día"
    assert decode_text("cañón,día".encode("latin-1")) == "cañón,día"
